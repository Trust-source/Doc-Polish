from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

GUIDE_TEMPLATE = {
    "title":   {"size": 20, "bold": True,  "align": "center", "font": "Times New Roman"},
    "heading": {"size": 13, "bold": True,  "align": "left", "font": "Times New Roman"},
    "bullet":  {"size": 10, "bold": False, "align": "left", "font": "Times New Roman"},
    "body":    {"size": 10, "bold": False, "align": "left", "font": "Times New Roman"},
    "blank":   {"size": 10, "bold": False, "align": "left", "font": "Times New Roman"},
}


ALIGN_MAP = {
    "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
    "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
    "right": WD_PARAGRAPH_ALIGNMENT.RIGHT
}


def extract_paragraphs(docx_path):
    doc = Document(docx_path)
    paragraphs = []

    for para in doc.paragraphs:

        if para.runs:
            first_run = para.runs[0]
            bold = first_run.bold
            if first_run.font.size:
                font_size = first_run.font.size.pt
        paragraphs.append({
            "text": para.text,
            "style": para.style.name,
            "bold": bold,
            "font_size": font_size
        })

    return paragraphs

def classify_role(para, index):
    style = para["style"]

    if style == "Title":
        return "title"
    if style in ("Heading 1", "Heading 2"):
        return "heading"
    if style == "List Bullet":
        return "bullet"
    if para["text"].strip() == "":
        return "blank"

    if para["bold"] and index < 5:
        return "heading"

    return "body"


def apply_style(paragraph, role, template):
    style = template[role]
    paragraph.alignment = ALIGN_MAP[style["align"]]
    for run in paragraph.runs:
        run.bold = style["bold"]
        run.font.size = Pt(style["size"])
        run.font.name = style["font"]

def format_document(input_path, output_path, template):
    doc = Document(input_path)
    paragraphs_data = extract_paragraphs(input_path)

    for index, (paragraph, para_data) in enumerate(zip(doc.paragraphs, paragraphs_data)):
        role = classify_role(para_data, index)
        apply_style(paragraph, role, template)

    doc.save(output_path)


format_document("sample.docx", "formatted_output.docx", GUIDE_TEMPLATE)